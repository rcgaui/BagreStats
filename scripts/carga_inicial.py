"""Carrega PELADA.xlsx para dentro do banco. Roda uma vez e pode ser repetido
sem duplicar nada.

Nao e um importador generico: a planilha tem uma quinta e 18 linhas, todas
limpas. Escrever tratamento de formato bagunçado aqui seria trabalho para um
problema que nao existe.
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.db import Session  # noqa: E402
from app.models import Cor, Jogador, Participacao, Partida  # noqa: E402

PLANILHA = Path(__file__).resolve().parent.parent / "PELADA.xlsx"

CORES_PADRAO = [
    ("PRETO", "#1a1a1a"),
    ("BRANCO", "#e8e8e8"),
    ("VERMELHO", "#d32f2f"),
    ("AZUL", "#1976d2"),
    ("VERDE", "#388e3c"),
    ("AMARELO", "#fbc02d"),
    ("LARANJA", "#f57c00"),
]


def _pegar_ou_criar(sessao, modelo, **campos):
    achado = sessao.query(modelo).filter_by(**campos).one_or_none()
    if achado:
        return achado, False
    novo = modelo(**campos)
    sessao.add(novo)
    sessao.flush()
    return novo, True


def main() -> None:
    wb = openpyxl.load_workbook(PLANILHA, data_only=True)
    partidas = list(wb["Partidas"].iter_rows(min_row=2, values_only=True))
    times = list(wb["Times"].iter_rows(min_row=2, values_only=True))

    with Session() as sessao:
        cores = {}
        for nome, codigo in CORES_PADRAO:
            cor, nova = _pegar_ou_criar(sessao, Cor, nome=nome)
            cor.codigo = codigo
            cores[nome] = cor
        print(f"cores disponiveis: {', '.join(cores)}")

        criados = 0
        for _, apelido, _ in times:
            if apelido and _pegar_ou_criar(sessao, Jogador, apelido=str(apelido).strip())[1]:
                criados += 1
        print(f"jogadores: {criados} cadastrados")

        for data, vencedor in partidas:
            if not data:
                continue
            partida, nova = _pegar_ou_criar(sessao, Partida, data=data.date())
            vencedor = (vencedor or "").strip().upper()
            partida.cor_vencedora_id = cores[vencedor].id if vencedor in cores else None

            for linha_data, apelido, time in times:
                if not linha_data or linha_data.date() != data.date():
                    continue
                jogador = sessao.query(Jogador).filter_by(apelido=str(apelido).strip()).one()
                _pegar_ou_criar(
                    sessao,
                    Participacao,
                    partida_id=partida.id,
                    jogador_id=jogador.id,
                    cor_id=cores[str(time).strip().upper()].id,
                )
            quantos = len(partida.participacoes)
            print(f"partida {partida.data}: vencedor {vencedor or 'EMPATE'}, {quantos} jogadores")

        sessao.commit()


if __name__ == "__main__":
    main()
